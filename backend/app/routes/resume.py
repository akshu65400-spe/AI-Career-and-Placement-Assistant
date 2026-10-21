from fastapi import APIRouter, UploadFile, File, Form, HTTPException
import os
import shutil
import re
from docx import Document
from pypdf import PdfReader


router = APIRouter(
    prefix="/resume",
    tags=["Resume"]
)


UPLOAD_DIR = "uploads/resumes"

os.makedirs(UPLOAD_DIR, exist_ok=True)


def extract_resume_text(file_path: str, extension: str) -> str:
    """Extract readable text from supported resume formats."""
    if extension == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(file_path).pages)

    if extension == ".docx":
        document = Document(file_path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_text = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + table_text)

    # Legacy .doc is a binary format. Recover readable text runs when possible;
    # DOCX or text-based PDF is preferred for more reliable analysis.
    with open(file_path, "rb") as resume_file:
        data = resume_file.read()
    runs = re.findall(rb"[\x20-\x7e]{4,}", data)
    text = "\n".join(run.decode("latin-1", errors="ignore") for run in runs)
    if not text.strip():
        text = data.decode("utf-16le", errors="ignore")
        text = re.sub(r"[^\x20-\x7e\r\n]+", " ", text)
    return text


def analyze_resume_text(text: str) -> dict:
    """Score resume content with a transparent, deterministic rubric."""
    normalized_text = re.sub(r"\s+", " ", text).strip()
    lower_text = normalized_text.lower()
    if len(normalized_text) < 40:
        raise HTTPException(
            status_code=400,
            detail="No readable resume text was found. Please upload a text-based PDF or DOCX (scanned PDFs need OCR).",
        )

    skills_database = [
        "python", "java", "javascript", "typescript", "react", "node.js",
        "fastapi", "sql", "mysql", "postgresql", "mongodb", "html", "css",
        "tailwind", "c++", "machine learning", "data science", "pandas",
        "numpy", "scikit-learn", "tensorflow", "docker", "git", "github",
    ]
    detected_skills = [
        skill.title() if skill not in {"c++", "sql", "html", "css", "git"} else skill.upper()
        for skill in skills_database
        if re.search(r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])", lower_text)
    ]

    section_patterns = {
        "summary": r"\b(summary|objective|professional profile|about me)\b",
        "education": r"\b(education|academic background|qualifications)\b",
        "experience": r"\b(experience|employment|work history|professional experience)\b",
        "projects": r"\b(projects|personal projects|academic projects)\b",
        "skills": r"\b(skills|technical skills|technologies)\b",
    }
    section_points = {"summary": 4, "education": 4, "experience": 7, "projects": 6, "skills": 4}
    sections_found = [name for name, pattern in section_patterns.items() if re.search(pattern, lower_text)]
    section_score = sum(section_points[name] for name in sections_found)
    skill_score = min(30, len(detected_skills) * 2)

    metric_count = len(re.findall(
        r"\b\d+(?:[.,]\d+)?\s*(?:%|percent\b|k\b|m\b|million\b|billion\b|users\b|customers\b|projects\b|years?\b|months?\b|hours?\b|days?\b|clients\b|revenue\b)",
        lower_text,
    ))
    achievement_score = min(15, metric_count * 3)
    action_verb_count = len(re.findall(
        r"\b(achieved|automated|built|created|delivered|designed|developed|improved|increased|launched|led|managed|optimized|reduced|implemented)\b",
        lower_text,
    ))
    action_score = min(10, action_verb_count * 2)

    word_count = len(re.findall(r"\b[\w+#.]+\b", normalized_text))
    length_score = 8 if 200 <= word_count <= 1500 else min(8, round(word_count / 25))
    contact_score = 4 * bool(re.search(r"\b[\w.+-]+@[\w.-]+\.[a-z]{2,}\b", lower_text))
    contact_score += 3 * bool(re.search(r"(?:\+?\d[\d\s().-]{7,}\d)", normalized_text))
    bullet_count = len(re.findall(r"(?:^|\n)\s*[-•*▪]", text))
    formatting_score = 5 if bullet_count >= 3 else min(5, bullet_count * 2)

    score_breakdown = {
        "sections": section_score,
        "skills": skill_score,
        "quantified_achievements": achievement_score,
        "action_verbs": action_score,
        "length": length_score,
        "contact_information": contact_score,
        "formatting": formatting_score,
    }
    score = min(100, sum(score_breakdown.values()))

    focus_skills = ["Machine Learning", "Docker", "FastAPI"]
    normalized_skills = {skill.lower() for skill in detected_skills}
    missing_skills = [skill for skill in focus_skills if skill.lower() not in normalized_skills]

    suggestions = []
    if metric_count < 2:
        suggestions.append("Add measurable achievements, such as time saved, users supported, or results improved.")
    if len(detected_skills) < 6:
        suggestions.append("Include more technical skills that are relevant to your target job.")
    if "projects" not in sections_found:
        suggestions.append("Add a projects section with the technologies used and the outcomes achieved.")
    if "experience" not in sections_found:
        suggestions.append("Add relevant work, internship, volunteer, or leadership experience.")
    if word_count < 200:
        suggestions.append("Add relevant detail so the resume gives a complete picture of your experience.")
    elif word_count > 1500:
        suggestions.append("Keep the resume concise by focusing on the most relevant experience and achievements.")
    if not suggestions:
        suggestions.append("Tailor your strongest achievements and skills to each job description.")

    return {
        "score": score,
        "skills": detected_skills,
        "missing": missing_skills,
        "suggestions": suggestions,
        "score_breakdown": score_breakdown,
    }


# ============================================================
# RESUME ANALYZE
# ============================================================

@router.post("/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    user_id: int | None = Form(None)
):

    # --------------------------------------------------------
    # Check file
    # --------------------------------------------------------

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected."
        )

    allowed_extensions = [
        ".pdf",
        ".doc",
        ".docx"
    ]

    filename_lower = file.filename.lower()

    if not any(
        filename_lower.endswith(ext)
        for ext in allowed_extensions
    ):
        raise HTTPException(
            status_code=400,
            detail="Only PDF, DOC and DOCX files are allowed."
        )

    # --------------------------------------------------------
    # Save file
    # --------------------------------------------------------

    safe_filename = re.sub(
        r"[^a-zA-Z0-9._-]",
        "_",
        file.filename
    )

    file_path = os.path.join(
        UPLOAD_DIR,
        safe_filename
    )

    try:

        with open(file_path, "wb") as buffer:

            shutil.copyfileobj(
                file.file,
                buffer
            )

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=f"Could not save resume: {str(e)}"
        )

    extension = os.path.splitext(filename_lower)[1]
    try:
        resume_text = extract_resume_text(file_path, extension)
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read resume contents. Please check the file or upload a text-based PDF/DOCX. ({str(e)})",
        )

    analysis = analyze_resume_text(resume_text)

    # --------------------------------------------------------
    # Response
    # --------------------------------------------------------

    # The frontend stores this result for the signed-in user's interview page.
    # No second upload is needed to start written or mock interview practice.
    return {
        "message": "Resume analyzed successfully",
        "filename": file.filename,
        "user_id": user_id,
        "resume_text": resume_text[:30000],
        **analysis,
    }