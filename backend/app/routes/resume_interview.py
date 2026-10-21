
import re
import random
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document
from io import BytesIO

router = APIRouter(
    prefix="/interview",
    tags=["Resume-Based Interview"],
)

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB

KNOWN_SKILLS = [
    "Python", "Java", "C++", "JavaScript", "TypeScript",
    "React", "Node.js", "Express", "FastAPI", "Flask",
    "HTML", "CSS", "MongoDB", "MySQL", "PostgreSQL",
    "SQL", "Machine Learning", "Deep Learning", "NLP",
    "Artificial Neural Network", "Random Forest", "Pandas",
    "NumPy", "Scikit-learn", "TensorFlow", "Keras",
    "Git", "GitHub", "Docker", "REST API", "Data Structures",
    "Algorithms", "OOP", "DBMS",
]


def extract_resume_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()

    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDFs are not supported.",
                )

            text = "\n".join(page.extract_text() or "" for page in reader.pages)

        elif suffix == ".docx":
            document = Document(BytesIO(content))
            paragraphs = [p.text for p in document.paragraphs]

            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(paragraphs)

        else:
            raise HTTPException(
                status_code=400,
                detail="Upload a PDF or DOCX resume.",
            )

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not read this resume. Try another PDF or DOCX file.",
        ) from exc

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if len(text) < 40:
        raise HTTPException(
            status_code=400,
            detail=(
                "Not enough readable text was found. "
                "If this is a scanned PDF, use a text-based PDF or DOCX."
            ),
        )

    return text[:30000]


def find_skills(text: str) -> list[str]:
    found = []

    for skill in KNOWN_SKILLS:
        pattern = re.escape(skill)
        if re.search(pattern, text, flags=re.IGNORECASE):
            found.append(skill)

    return found


def extract_projects(text: str) -> list[str]:
    """
    Look for text near common project headings.
    This is heuristic extraction, not full semantic understanding.
    """
    lines = [line.strip(" •-\t") for line in text.splitlines()]
    projects = []

    project_heading = re.compile(
        r"^(?:academic\s+|personal\s+)?projects?\s*:?$",
        re.IGNORECASE,
    )

    stop_heading = re.compile(
        r"^(skills|education|experience|certifications|achievements|"
        r"languages|interests|summary|career objective)\b",
        re.IGNORECASE,
    )

    in_projects = False

    for line in lines:
        if project_heading.search(line):
            in_projects = True
            continue

        if in_projects and stop_heading.search(line):
            break

        if in_projects and len(line) >= 5 and line.casefold() not in {item.casefold() for item in projects}:
            projects.append(line)

        if len(projects) >= 15:
            break

    return projects


def build_questions(
    resume_text: str,
    role: str,
    difficulty: str,
    interview_round: str = "Technical",
) -> tuple[list[str], list[str], list[str]]:
    skills = find_skills(resume_text)
    projects = extract_projects(resume_text)

    role = role.strip() or "Software Developer"

    round_key = interview_round.strip().lower()
    if round_key not in {"technical", "hr", "coding"}:
        round_key = "technical"

    if round_key == "hr":
        questions = [
            "Please introduce yourself and summarize your background from your resume.",
            f"Why are you interested in the {role} role, and what would you bring to the team?",
            "What is one strength you would bring to this role? Give a real example.",
            "What is a professional weakness you are actively working to improve?",
            "Tell me about a time you handled a disagreement with a teammate.",
            "Describe a time you received constructive feedback. What did you do with it?",
            "Tell me about a time you had to meet a difficult deadline.",
            "Describe a mistake or setback and what you learned from it.",
            "How do you prioritize when several important tasks are due at once?",
            "Tell me about a time you took initiative without being asked.",
            "How do you adapt when project requirements or priorities change?",
            "What kind of work environment helps you do your best work?",
            "Describe a time you helped a team reach a shared goal.",
            "What accomplishment from your resume are you most proud of, and why?",
            "What are your career goals for the next few years?",
            "Why should we select you for this position?",
            "How do you stay motivated when a task becomes challenging?",
            "Tell me about a time you had to learn something quickly to complete a task.",
        ]
        questions.extend(
            f"Your resume mentions '{project[:120]}'. What was your role, and what did you learn from working on it?"
            for project in projects
        )
        selected_questions = random.sample(list(dict.fromkeys(questions)), 15)
        random.shuffle(selected_questions)
        return selected_questions, skills, projects

    if round_key == "coding":
        coding_questions = [
            "Write a function that finds the first non-repeating character in a string. Explain its time and space complexity.",
            "Given an array of integers, find two values that add up to a target. Describe an efficient approach.",
            "How would you detect whether a linked list contains a cycle?",
            "Write logic to validate whether brackets in a string are balanced.",
            "Given a list of numbers, remove duplicates while preserving their original order.",
            "How would you find the longest substring without repeating characters?",
            "Explain how you would search for a value in a sorted array and state the complexity.",
            "Given a string, determine whether it is a palindrome while ignoring spaces and case.",
            "How would you find the most frequent element in an array?",
            "Write a function to merge two sorted arrays into one sorted result.",
            "How would you traverse a binary tree level by level?",
            "Given a list of intervals, merge all overlapping intervals.",
            "How would you find the missing number from a sequence containing values from 0 to n?",
            "Design a stack that can return its minimum value in constant time.",
            "How would you test your solution for edge cases and large inputs?",
            "Explain the time and space complexity of a recent algorithm or feature you implemented.",
            "Choose a project from your resume. What part would you implement first in code, and why?",
        ]
        resume_coding_questions = [
            f"Using {skill}, describe or write a small function that solves a practical problem. What edge cases would you test?"
            for skill in skills
        ]
        resume_coding_questions.extend(
            f"For the project '{project[:120]}' on your resume, describe one coding task you completed and how you would test it."
            for project in projects
        )
        coding_questions = list(dict.fromkeys(coding_questions))
        resume_coding_questions = list(dict.fromkeys(resume_coding_questions))
        selected_questions = random.sample(coding_questions, 10)
        selected_questions.extend(
            random.sample(resume_coding_questions, min(5, len(resume_coding_questions)))
        )
        remaining_count = 15 - len(selected_questions)
        if remaining_count:
            selected_questions.extend(random.sample(
                [question for question in coding_questions if question not in selected_questions],
                remaining_count,
            ))
        random.shuffle(selected_questions)
        return selected_questions, skills, projects

    general_questions = [
        "Please introduce yourself and summarize your background from your resume.",
        "Describe a difficult technical problem you faced and how you solved it.",
        "What would you improve in the project or experience described in your resume?",
        f"What makes your skills and experience suitable for a {role} position?",
        "What is one technical skill you are currently improving, and how?",
        "How do you prioritize your work when several tasks need attention?",
        "Describe a technical decision you made and the trade-offs you considered.",
        "Walk me through a project you are proud of and the impact it had.",
        "How do you test your work and make sure it meets the requirements?",
        "How do you collaborate with teammates when building a technical solution?",
        "How do you keep your code or project maintainable as it grows?",
        "Tell me about a time you had to adapt to a change in project requirements.",
        "What measurable result or outcome from your work are you most proud of?",
        "How do you learn an unfamiliar tool or technology when a project requires it?",
        "How do you verify that a solution is reliable before delivering it?",
    ]

    if difficulty.lower() in {"hard", "advanced"}:
        general_questions.extend([
            "How would you test, secure, and scale the system described in your resume?",
            "Describe a technical trade-off you made or would make in one of your projects.",
        ])

    personalized_questions = []
    for skill in skills:
        personalized_questions.extend([
            f"Your resume lists {skill}. Describe a task or project where you used it and what you built.",
            f"What challenge did you face while using {skill}, and how did you address it?",
            f"How would you explain the strengths and limitations of {skill} to a teammate?",
            f"How could {skill} work with another technology from your resume?",
        ])

    project_questions = []
    for project in projects:
        project_detail = project[:150]
        project_prompts = [
            f"Your resume mentions '{project_detail}'. What was your contribution, and what changed because of your work?",
            f"What technical challenge came up in '{project_detail}', and how did you resolve it?",
            f"Why did you choose the tools or design approach used in '{project_detail}'?",
            f"If you revisited '{project_detail}', what would you improve and why?",
        ]
        project_questions.append(random.choice(project_prompts))

    if len(skills) >= 2:
        skill_pairs = list(zip(skills, skills[1:]))
        personalized_questions.extend(
            f"Your resume lists {first_skill} and {second_skill}. How would you combine them to build a useful feature?"
            for first_skill, second_skill in skill_pairs
        )

    # Ask about every detected project, then fill and shuffle the remaining slots.
    personalized_questions = list(dict.fromkeys(personalized_questions))
    general_questions = list(dict.fromkeys(general_questions))
    selected_questions = project_questions[:15]
    remaining_pool = list(dict.fromkeys(personalized_questions + general_questions))
    remaining_pool = [question for question in remaining_pool if question not in selected_questions]
    selected_questions.extend(random.sample(remaining_pool, 15 - len(selected_questions)))
    random.shuffle(selected_questions)

    return selected_questions, skills, projects


@router.post("/mock/start-from-resume")
async def start_mock_interview_from_resume(
    role: str = Form("Software Developer"),
    difficulty: str = Form("Easy"),
    resume: UploadFile = File(...),
):
    filename = resume.filename or ""

    if Path(filename).suffix.lower() not in {".pdf", ".docx"}:
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX resumes are supported.",
        )

    content = await resume.read(MAX_FILE_SIZE + 1)

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="Resume file must be 5 MB or smaller.",
        )

    text = extract_resume_text(filename, content)

    questions, skills, projects = build_questions(
        text,
        role,
        difficulty,
    )

    return {
        "role": role,
        "difficulty": difficulty,
        "questions": questions,
        "question": questions[0],
        "total_questions": len(questions),
        "detected_skills": skills,
        "detected_projects": projects,
        "generation_method": "resume-based rules",
    }


class SavedResumeInterviewStart(BaseModel):
    resume_text: str = Field(min_length=40, max_length=30000)
    role: str = Field(default="Software Developer", min_length=1, max_length=120)
    difficulty: str = Field(default="Easy", max_length=30)
    interview_round: str = Field(default="Technical", max_length=30)


@router.post("/start-from-resume-data")
def start_interview_from_saved_resume(data: SavedResumeInterviewStart):
    """Generate questions from text extracted by Resume AI; no file upload here."""
    resume_text = data.resume_text.strip()
    if len(resume_text) < 40:
        raise HTTPException(
            status_code=400,
            detail="No analyzed resume text was found. Open Resume AI and analyze your resume first.",
        )

    questions, skills, projects = build_questions(
        resume_text, data.role, data.difficulty, data.interview_round
    )
    return {
        "role": data.role,
        "difficulty": data.difficulty,
        "interview_round": data.interview_round,
        "questions": questions,
        "question": questions[0],
        "total_questions": len(questions),
        "detected_skills": skills,
        "detected_projects": projects,
        "generation_method": "saved Resume AI analysis",
    }
